from copy import copy
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from bom_post_process.errors import BomError
from bom_post_process.soldering import gen_soldering
from bom_post_process.values import gen_value
from bom_post_process.workbook import AD_HEADERS, BOM_HEADERS, split_designators
from bom_post_process.standard import RESOURCES, Template, read_bom, read_exported_content


class SolderingTests(unittest.TestCase):
    def setUp(self):
        self.temporary = TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)
        self.source = self.folder / "样例_已解析.xlsx"
        self.output = self.folder / "焊接.xlsx"
        book = Workbook()
        original = book.active
        original.title = "AD"
        original.append(AD_HEADERS)
        for row in [
            ["0402-X5R104K250", "", "C1,C2", "C0402", "Cap", 2],
            ["0402K1003", "", "R1", "R0402", "Res", 1],
            ["0402K1003/NC", "", "R2,R3", "R0402", "Res", 2],
            ["0402-X5R102K250(NC)", "", "C3", "C0402", "Cap", 1],
            ["IC", "", "U1", "QFN", "IC", 1],
        ]:
            original.append(row)
        bom = book.create_sheet("BOM参数表")
        bom.append(BOM_HEADERS)
        for row in [
            [1, "PCB制板", "PCB制板", 2, None, "占位项"],
            [2, "0402-X5R104K250", "人工电容名称", 2, "C1,C2", "人工改为100nF/50V，指定供应商"],
            [3, "0402K1003", "人工电阻名称", 1, "R1", "人工修改：100k/0.1%"],
            [4, "0402K1003", "电阻", 2, "r2、R3", "未焊接人工备注"],
            [5, "0402-X5R102K250", "电容", 1, "C3", "NC人工补充"],
            [6, "IC", "自定义芯片", 1, "U1", "保留完整人工备注"],
            [7, "工装", "工装项", 1, None, "无位号保留"],
        ]:
            bom.append(row)
        bom["F3"].font = Font(name="宋体", size=12, bold=True)
        bom["F3"].fill = PatternFill("solid", fgColor="FFFF00")
        bom["F3"].alignment = Alignment(wrap_text=True)
        bom.row_dimensions[3].height = 48
        bom.column_dimensions["F"].width = 52
        bom.freeze_panes = "A2"
        book.save(self.source)
        book.close()

    def test_manual_values_copied_verbatim_with_styles_without_reparsing(self):
        before = self.source.read_bytes()
        report = gen_soldering(self.source, self.output)
        self.assertEqual(self.source.read_bytes(), before)
        self.assertEqual((report.written, report.excluded_nc, report.excluded_pcb, report.nc_designators), (4, 2, 1, 3))
        source, result = load_workbook(self.source), load_workbook(self.output)
        try:
            self.assertEqual(result.sheetnames, ["焊接表", "未焊接元件"])
            bom = source["BOM参数表"]
            for sheet_name, kept in [("焊接表", [3, 4, 7, 8]), ("未焊接元件", [5, 6])]:
                target = result[sheet_name]
                for output_row, input_row in enumerate([1, *kept], start=1):
                    for column in range(1, 7):
                        original, copied = bom.cell(input_row, column), target.cell(output_row, column)
                        self.assertEqual(original.value, copied.value)
                        for attribute in ("font", "fill", "alignment", "border", "number_format"):
                            self.assertEqual(copy(getattr(original, attribute)), copy(getattr(copied, attribute)))
                self.assertEqual(target.column_dimensions["F"].width, 52)
            self.assertEqual(result["焊接表"]["F3"].value, "人工修改：100k/0.1%")
            self.assertEqual(result["焊接表"].row_dimensions[2].height, 48)
        finally:
            source.close()
            result.close()

    def test_raw_ad_input_requires_manual_table(self):
        book = load_workbook(self.source)
        del book["BOM参数表"]
        book.save(self.source)
        book.close()
        with self.assertRaisesRegex(BomError, "先运行 gen_value 并人工修改备注"):
            gen_soldering(self.source, self.output)
        self.assertFalse(self.output.exists())

    def test_explicit_sheet_selects_reviewed_table(self):
        book = load_workbook(self.source)
        other = book.copy_worksheet(book["BOM参数表"])
        other.title = "另一份"
        book.save(self.source)
        book.close()
        with self.assertRaisesRegex(BomError, "多个整理后"):
            gen_soldering(self.source, self.output)
        report = gen_soldering(self.source, self.output, sheet="BOM参数表")
        self.assertEqual(report.bom_sheet, "BOM参数表")

    def test_mixed_nc_row_rejected_without_output(self):
        book = load_workbook(self.source)
        book["BOM参数表"]["E5"] = "R1, R2"
        book.save(self.source)
        book.close()
        with self.assertRaisesRegex(BomError, "混有贴装和不贴装.*R2"):
            gen_soldering(self.source, self.output)
        self.assertFalse(self.output.exists())

    def test_multiple_original_ad_sheets_and_date_epoch_are_preserved(self):
        book = load_workbook(self.source)
        book.copy_worksheet(book["AD"]).title = "另一原始表"
        book.epoch = CALENDAR_MAC_1904
        book.save(self.source)
        book.close()
        with self.assertRaisesRegex(BomError, "多个 AD"):
            gen_soldering(self.source, self.output)
        report = gen_soldering(self.source, self.output, sheet="BOM参数表", ad_sheet="AD")
        self.assertEqual((report.source_sheet, report.written, report.excluded_nc), ("AD", 4, 2))
        result = load_workbook(self.output)
        try:
            self.assertEqual(result.epoch, CALENDAR_MAC_1904)
        finally:
            result.close()

    def test_empty_nc_has_header_and_default_name(self):
        book = load_workbook(self.source)
        book["AD"]["A4"] = "0402K1003"
        book["AD"]["A5"] = "0402-X5R102K250"
        book.save(self.source)
        book.close()
        report = gen_soldering(self.source)
        self.assertEqual(report.output.name, "样例_焊接表.xlsx")
        self.assertEqual((report.written, report.excluded_nc), (6, 0))
        book = load_workbook(report.output)
        try:
            self.assertEqual(book["未焊接元件"].max_row, 1)
        finally:
            book.close()

    def test_invalid_data_and_layout_rejected(self):
        for mode in ("formula", "merged", "nc_empty", "excel_error", "empty_bom"):
            with self.subTest(mode=mode):
                book = load_workbook(self.source)
                if mode == "formula":
                    book["BOM参数表"]["F3"] = "=A3"
                elif mode == "merged":
                    book["BOM参数表"].merge_cells("E3:F3")
                elif mode == "nc_empty":
                    book["AD"]["C4"] = None
                elif mode == "excel_error":
                    book["BOM参数表"]["F3"] = "#REF!"
                else:
                    book["BOM参数表"].delete_rows(2, 20)
                invalid = self.folder / f"{mode}.xlsx"
                book.save(invalid)
                book.close()
                with self.assertRaises(BomError):
                    gen_soldering(invalid, self.output)
                self.assertFalse(self.output.exists())

    def test_output_and_source_never_overwritten(self):
        self.output.write_bytes(b"keep existing")
        for output in (self.source, self.output):
            before = output.read_bytes()
            with self.assertRaises(BomError):
                gen_soldering(self.source, output)
            self.assertEqual(output.read_bytes(), before)

    def test_real_input_and_manual_change_survive_two_stages(self):
        source = Path(__file__).resolve().parents[1] / "input" / "DAQ_1010A_V1R0.xlsx"
        parsed = self.folder / "parsed.xlsx"
        gen_value(source, parsed)
        book = load_workbook(parsed)
        book.active["G63"] = "人工确认100k/0.1%，勿重新解析"
        book.save(parsed)
        book.close()
        report = gen_soldering(parsed, self.output)
        self.assertEqual((report.written, report.excluded_nc, report.nc_designators), (134, 8, 29))
        book = load_workbook(self.output)
        try:
            values = list(book["焊接表"].values)[1:]
            reviewed = [row for row in values if "R1" in split_designators(row[4])]
            self.assertEqual(reviewed[0][5], "人工确认100k/0.1%，勿重新解析")
            self.assertEqual(sum(row[3] for row in list(book["未焊接元件"].values)[1:]), 29)
        finally:
            book.close()

    def test_reordered_note_input_carries_manual_ic_selection_through_to_word(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "已确认"
        sheet.append(["note", "Maker", "Quantity", "LibRef", "Comment", "Designator", "Description", "Footprint"])
        note = "  手工选型说明：指定厂家 & 型号\n工作温度 -40~85℃  "
        sheet.append([note, "厂家A", 1, "IC", "ChipX", "U1", "不要复制此描述到备注", "QFN"])
        sheet.append(["未焊接人工备注", "厂家B", 2, "Res", "0402K1003/NC", "R1, R2", "电阻描述", "R0402"])
        sheet["A2"].font = Font(name="宋体", size=12, bold=True)
        sheet["A2"].fill = PatternFill("solid", fgColor="FFFF00")
        book.save(self.source)
        book.close()
        before = self.source.read_bytes()
        report = gen_soldering(self.source, self.output)
        self.assertEqual((report.written, report.excluded_nc, report.bom_sheet), (1, 1, "已确认"))
        self.assertEqual(self.source.read_bytes(), before)
        exported = load_workbook(self.output)
        self.assertEqual(exported["焊接表"]["F2"].value, note)
        self.assertTrue(exported["焊接表"]["F2"].font.bold)
        self.assertEqual(exported["未焊接元件"]["F2"].value, "未焊接人工备注")
        exported.close()
        _, records = read_bom(self.output)
        self.assertEqual(records, [["1", "ChipX", "IC", "1", "U1", note]])
        template = Template(RESOURCES / "bom_template.docx")
        word = self.folder / "reviewed.docx"
        template.write(word, [records])
        self.assertEqual(read_exported_content(word, template)[0], records)

    def test_formula_in_note_is_rejected_without_export(self):
        book = Workbook()
        book.active.append([*AD_HEADERS, "note"])
        book.active.append(["IC", "", "U1", "QFN", "IC", 1, "=1+1"])
        book.save(self.source)
        book.close()
        with self.assertRaisesRegex(BomError, "note.*公式"):
            gen_soldering(self.source, self.output)
        self.assertFalse(self.output.exists())


if __name__ == "__main__":
    unittest.main()
