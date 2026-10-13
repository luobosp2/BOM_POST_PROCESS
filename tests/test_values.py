from copy import copy
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from bom_post_process.errors import BomError
from bom_post_process.values import gen_value
from bom_post_process.workbook import AD_HEADERS, BOM_HEADERS, read_ad_workbook, tokenize_designators


class ValuesTests(unittest.TestCase):
    def setUp(self):
        self.temporary = TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)
        self.source = self.folder / "样例.xlsx"
        self.output = self.folder / "已解析.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "AD导出BOM"
        sheet.append(AD_HEADERS)
        for row in [
            ["0402-X5R104K250", "原电容描述", "C1, C2", "C 0402L", "Cap", 2],
            ["0402K1003", "原电阻描述", "R1", "RES-0402", "Res1", 1],
            ["0402K1003/NC", "", "R2, R3", "0402R", "Res", 2],
            ["0402-X5R102K250(NC)", "", "C3", "C0402", "Cap", 1],
            ["1206B104K101NT(FH)", "人工原描述", "C4", "C1206", "Cap", 1],
            ["M1206", "降压芯片", "U1", "M1206", "M1206", 1],
        ]:
            sheet.append(row)
        sheet["B2"].font = Font(name="宋体", bold=True)
        book.save(self.source)
        book.close()

    def test_appends_review_columns_without_changing_original_cells_or_sheets(self):
        before = self.source.read_bytes()
        report = gen_value(self.source, self.output)
        self.assertEqual(self.source.read_bytes(), before)
        self.assertEqual((report.capacitors, report.resistors, report.parsed), (2, 2, 4))
        self.assertEqual((report.total, report.skipped, report.sheet), (6, 1, "AD导出BOM"))
        self.assertEqual([(issue.row, issue.model) for issue in report.issues], [(6, "1206B104K101NT(FH)")])
        original, book = load_workbook(self.source), load_workbook(self.output)
        try:
            self.assertEqual(book.sheetnames, original.sheetnames)
            self.assertEqual(list(book.active.iter_rows(max_col=6, values_only=True)), list(original.active.values))
            self.assertEqual(copy(book["AD导出BOM"]["B2"].font), copy(original.active["B2"].font))
            bom = book.active
            self.assertEqual(tuple(cell.value for cell in bom[1]), (*AD_HEADERS, "name", "note", "not_bought", "not_welded"))
            self.assertEqual(bom["G2"].value, "贴片电容")
            self.assertEqual(bom["G3"].value, "贴片电阻")
            self.assertEqual(bom["G7"].value, "M1206")
            self.assertEqual(bom["H2"].value, "100nF/25V/0402/X5R/10%")
            self.assertEqual(bom["H3"].value, "100k/0402/1%")
            self.assertEqual(bom["A4"].value, "0402K1003/NC")
            self.assertEqual(bom["H5"].value, "1nF/25V/0402/X5R/10%")
            self.assertIn(bom["H6"].value, (None, ""))
            self.assertIn(bom["H7"].value, (None, ""))
            self.assertEqual([bom.cell(row, 9).value for row in range(2, 8)], [""] * 6)
            self.assertEqual([bom.cell(row, 10).value for row in range(2, 8)], ["", "", "NC", "NC", "", ""])
            self.assertEqual(bom["A7"].value, "M1206")
            self.assertEqual(book.active.title, original.active.title)
        finally:
            original.close()
            book.close()

    def test_other_sheets_and_original_layout_are_not_changed(self):
        book = load_workbook(self.source)
        bom = book.create_sheet("整理后BOM")
        bom.append(BOM_HEADERS)
        bom.append([1, "PCB制板", "PCB制板", 2, None, "保留占位项"])
        bom.append([2, "0402K1003", "人工电阻名", 1, "R1", "待解析"])
        bom.append([3, "0402K1003/NC", "贴片电阻", 2, "R2,R3", "待解析"])
        bom.append([4, "1206B104K101NT(FH)", "贴片电容", 1, "C4", "厂商参数，保留"])
        bom.append([5, "M1206", "人工芯片名", 1, "U1", "人工补充备注"])
        bom["F5"].font = Font(name="宋体", size=12, bold=True)
        bom["F5"].fill = PatternFill("solid", fgColor="FFFF00")
        bom["F5"].alignment = Alignment(wrap_text=True)
        bom.row_dimensions[5].height = 48
        bom.column_dimensions["F"].width = 52
        book.save(self.source)
        book.close()
        report = gen_value(self.source, self.output)
        self.assertEqual((report.total, report.resistors, report.sheet), (6, 2, "AD导出BOM"))
        source, result = load_workbook(self.source), load_workbook(self.output)
        try:
            before, after = source["整理后BOM"], result["整理后BOM"]
            self.assertEqual(after["B2"].value, "PCB制板")
            self.assertEqual(after["C3"].value, "人工电阻名")
            self.assertEqual(after["F3"].value, "待解析")
            self.assertEqual(after["B4"].value, "0402K1003/NC")
            self.assertEqual(after["F5"].value, "厂商参数，保留")
            self.assertEqual(after["F6"].value, "人工补充备注")
            for attribute in ("font", "fill", "alignment"):
                self.assertEqual(copy(getattr(before["F5"], attribute)), copy(getattr(after["F5"], attribute)))
            self.assertEqual(after.row_dimensions[5].height, 48)
            self.assertEqual(after.column_dimensions["F"].width, 52)
            self.assertEqual(list(before.values), list(after.values))
        finally:
            source.close()
            result.close()

    def test_default_name_and_existing_files_are_not_overwritten(self):
        report = gen_value(self.source)
        self.assertEqual(report.output.name, "样例_已解析.xlsx")
        for output in (self.source, report.output):
            before = output.read_bytes()
            with self.assertRaises(BomError):
                gen_value(self.source, output)
            self.assertEqual(output.read_bytes(), before)

    def test_formula_and_excel_errors_do_not_create_output(self):
        for cell, value, expected in [("F2", "=1+1", "公式"), ("A2", "#REF!", "Excel 错误")]:
            with self.subTest(cell=cell):
                book = load_workbook(self.source)
                book.active[cell] = value
                bad = self.folder / f"bad_{cell}.xlsx"
                book.save(bad)
                book.close()
                with self.assertRaisesRegex(BomError, expected):
                    gen_value(bad, self.output)
                self.assertFalse(self.output.exists())

    def test_empty_table_is_rejected(self):
        book = Workbook()
        book.active.append(AD_HEADERS)
        book.save(self.source)
        book.close()
        with self.assertRaisesRegex(BomError, "没有物料数据"):
            gen_value(self.source, self.output)
        self.assertFalse(self.output.exists())

    def test_formula_like_text_stays_text(self):
        book = load_workbook(self.source)
        book.active["A7"] = "=not_a_formula"
        book.active["A7"].data_type = "s"
        book.save(self.source)
        book.close()
        gen_value(self.source, self.output)
        book = load_workbook(self.output)
        try:
            self.assertEqual(book.active["A7"].value, "=not_a_formula")
            self.assertEqual(book.active["A7"].data_type, "s")
        finally:
            book.close()

    def test_real_input_regression(self):
        source = Path(__file__).resolve().parents[1] / "input" / "DAQ_1010A_V1R0.xlsx"
        before = source.read_bytes()
        report = gen_value(source, self.output)
        self.assertEqual((report.capacitors, report.resistors, report.skipped), (34, 57, 49))
        self.assertEqual(report.total, 142)
        self.assertEqual([issue.row for issue in report.issues], [11, 43])
        self.assertEqual(source.read_bytes(), before)
        book = load_workbook(self.output)
        try:
            rows = list(book.active.values)[1:]
            self.assertEqual(len(rows), 142)
            self.assertEqual(sum(row[5] for row in rows), sum(row.quantity for row in read_ad_workbook(source).rows))
            self.assertEqual(sum(bool(row[7]) for row in rows), 91)
            self.assertEqual(sum(bool(row[9]) for row in rows), 8)
            self.assertTrue(all(row[8] == "" for row in rows))
        finally:
            book.close()

    def test_only_target_worksheet_zip_member_changes(self):
        book = load_workbook(self.source)
        sheet = book.active
        sheet.column_dimensions["A"].width = 28
        sheet.row_dimensions[2].height = 39
        sheet.freeze_panes = "C3"
        sheet.auto_filter.ref = "A1:F7"
        sheet.print_area = "A1:F7"
        sheet.page_setup.orientation = "landscape"
        sheet["H1"] = "Extra"
        sheet["H2"] = "=F2+1"
        sheet.merge_cells("H6:I6")
        book.create_sheet("Cover")["A1"] = "其他页面保持原样"
        book.active = 1
        book.save(self.source)
        book.close()
        with ZipFile(self.source, "a") as archive:
            archive.writestr("customXml/retained.xml", b"<original>opaque metadata</original>")
        gen_value(self.source, self.output)
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        with ZipFile(self.source) as original, ZipFile(self.output) as result:
            self.assertEqual(original.namelist(), result.namelist())
            for name in original.namelist():
                if name != "xl/worksheets/sheet1.xml":
                    self.assertEqual(original.read(name), result.read(name), name)
            before = ET.fromstring(original.read("xl/worksheets/sheet1.xml"))
            after = ET.fromstring(result.read("xl/worksheets/sheet1.xml"))
            for cell in before.findall("s:sheetData/s:row/s:c", ns):
                copied = after.find(f"s:sheetData/s:row/s:c[@r='{cell.get('r')}']", ns)
                self.assertEqual(ET.tostring(cell), ET.tostring(copied))
            for name in ("cols", "sheetViews", "sheetFormatPr", "pageMargins", "pageSetup", "autoFilter", "mergeCells"):
                node = before.find(f"s:{name}", ns)
                copied = after.find(f"s:{name}", ns)
                if name == "cols":
                    for old, new in zip(node, copied):
                        self.assertEqual(ET.tostring(old), ET.tostring(new))
                elif node is not None:
                    self.assertEqual(ET.tostring(node), ET.tostring(copied), name)
        result = load_workbook(self.output)
        self.assertEqual(result.active.title, "Cover")
        self.assertEqual([result["AD导出BOM"].cell(1, column).value for column in range(10, 14)],
                         ["name", "note", "not_bought", "not_welded"])
        result.close()

    def test_legacy_last_note_migrates_with_content_style_and_width(self):
        book = load_workbook(self.source)
        sheet = book.active
        sheet["G1"] = " Note "
        sheet["G2"] = "人工选型：50V\n指定厂家，保留原文  "
        sheet["G6"] = "厂商电容，已人工确认"
        sheet["G7"] = "芯片选型说明"
        sheet["G2"].font = Font(name="宋体", bold=True)
        sheet["G2"].alignment = Alignment(wrap_text=True)
        sheet.column_dimensions["G"].width = 51
        sheet.print_area = "A1:F7"
        book.create_sheet("Cover")["A1"] = "原封不动"
        book.save(self.source)
        book.close()
        report = gen_value(self.source, self.output)
        self.assertEqual((report.capacitors, report.resistors), (1, 2))
        result = load_workbook(self.output)
        try:
            self.assertEqual(result.active.max_column, 10)
            self.assertEqual(result.active["H1"].value, " Note ")
            self.assertEqual(result.active["H2"].value, "人工选型：50V\n指定厂家，保留原文  ")
            self.assertEqual(result.active["H3"].value, "100k/0402/1%")
            self.assertEqual(result.active["H6"].value, "厂商电容，已人工确认")
            self.assertEqual(result.active["H7"].value, "芯片选型说明")
            self.assertEqual(result.active["G2"].value, "贴片电容")
            self.assertEqual(result.active["J4"].value, "NC")
            self.assertEqual(result.active.column_dimensions["H"].width, 51)
            self.assertEqual(result.active["H2"].font.name, "宋体")
            self.assertTrue(result.active["H2"].font.bold)
            self.assertTrue(result.active["H2"].alignment.wrap_text)
        finally:
            result.close()
        with ZipFile(self.source) as original, ZipFile(self.output) as result:
            for name in original.namelist():
                if name != "xl/worksheets/sheet1.xml":
                    self.assertEqual(original.read(name), result.read(name), name)

    def test_existing_review_fields_and_manual_nc_override_are_preserved(self):
        gen_value(self.source, self.output)
        book = load_workbook(self.output)
        sheet = book.active
        sheet["G2"] = "人工电容名称  "
        sheet["G3"] = None
        sheet["H2"] = "人工参数\n保留尾部空格  "
        sheet["H6"] = "手动确认"
        sheet["H7"] = "自选芯片"
        sheet["I2"] = "不要采购"
        sheet["J3"] = "仅采购"
        sheet["J4"] = None
        sheet["J5"] = "自定义原因"
        book.save(self.output)
        book.close()
        output = self.folder / "再次解析.xlsx"
        before = self.output.read_bytes()
        report = gen_value(self.output, output)
        self.assertEqual((report.parsed, report.skipped), (0, 6))
        self.assertEqual(output.read_bytes(), before)
        result = load_workbook(output)
        try:
            self.assertEqual(result.active["G2"].value, "人工电容名称  ")
            self.assertIsNone(result.active["G3"].value)
            self.assertEqual(result.active["H2"].value, "人工参数\n保留尾部空格  ")
            self.assertEqual(result.active["I2"].value, "不要采购")
            self.assertEqual(result.active["J3"].value, "仅采购")
            self.assertIsNone(result.active["J4"].value)
            self.assertEqual(result.active["J5"].value, "自定义原因")
        finally:
            result.close()

    def test_nc_suffix_matching_is_case_insensitive_for_new_flags(self):
        book = load_workbook(self.source)
        book.active["A4"] = "0402K1003/nc"
        book.active["A5"] = "0402-X5R102K250(nC)  "
        book.save(self.source)
        book.close()
        gen_value(self.source, self.output)
        result = load_workbook(self.output)
        try:
            self.assertEqual(result.active["J4"].value, "NC")
            self.assertEqual(result.active["J5"].value, "NC")
            self.assertEqual(result.active["H4"].value, "100k/0402/1%")
        finally:
            result.close()

    def test_unsafe_legacy_note_migrations_do_not_create_output(self):
        for complication in ("following_column", "formula_reference", "filter_reference"):
            with self.subTest(complication=complication):
                book = load_workbook(self.source)
                sheet = book.active
                sheet["G1"] = "note"
                sheet["G2"] = "人工备注"
                if complication == "following_column":
                    sheet["H1"] = "Extra"
                elif complication == "formula_reference":
                    book.create_sheet("Reference")["A1"] = "='AD导出BOM'!G2"
                else:
                    sheet.auto_filter.ref = "A1:G7"
                source = self.folder / (complication + ".xlsx")
                book.save(source)
                book.close()
                with self.assertRaisesRegex(BomError, "旧版 note"):
                    gen_value(source, self.output)
                self.assertFalse(self.output.exists())

    def test_reordered_columns_and_extra_fields_keep_their_positions(self):
        book = Workbook()
        book.active.append(["Maker", "Quantity", "LIBREF", "Footprint", "Designator", "Description", " comment "])
        book.active.append(["厂家A", 1, "Res", "R0402", "R1", "源描述", "0402K1003"])
        book.save(self.source)
        book.close()
        gen_value(self.source, self.output)
        book = load_workbook(self.output)
        self.assertEqual(list(book.active.values), [
            ("Maker", "Quantity", "LIBREF", "Footprint", "Designator", "Description", " comment ", "name", "note", "not_bought", "not_welded"),
            ("厂家A", 1, "Res", "R0402", "R1", "源描述", "0402K1003", "贴片电阻", "100k/0402/1%", "", ""),
        ])
        book.close()


class WorkbookReaderTests(unittest.TestCase):
    def test_reordered_headers_case_and_explicit_sheet(self):
        with TemporaryDirectory() as folder:
            source = Path(folder) / "input.xlsx"
            book = Workbook()
            book.active.title = "说明"
            book.active.append(["说明"])
            sheet = book.create_sheet("AD")
            sheet.append([" quantity ", "LIBREF", "Footprint", "Designator", "Description", "Comment"])
            sheet.append([1, "Res1", "R0402", "R1", "desc", "0402K1003"])
            book.save(source)
            book.close()
            table = read_ad_workbook(source)
            self.assertEqual((table.sheet, len(table.rows)), ("AD", 1))
            self.assertEqual((table.rows[0].comment, table.rows[0].quantity), ("0402K1003", 1))
            self.assertEqual(read_ad_workbook(source, sheet="AD"), table)
            with self.assertRaisesRegex(BomError, "缺少 AD 表头"):
                read_ad_workbook(source, sheet="说明")
            with self.assertRaisesRegex(BomError, "找不到工作表"):
                read_ad_workbook(source, sheet="missing")

    def test_multiple_ad_sheets_require_selection(self):
        with TemporaryDirectory() as folder:
            source = Path(folder) / "input.xlsx"
            book = Workbook()
            book.active.append(AD_HEADERS)
            book.create_sheet("AD2").append(AD_HEADERS)
            book.save(source)
            book.close()
            with self.assertRaisesRegex(BomError, "多个 AD"):
                read_ad_workbook(source)
            self.assertEqual(read_ad_workbook(source, "AD2").rows, ())

    def test_tokenizer_preserves_case_and_duplicates(self):
        self.assertEqual(tokenize_designators("r1、R2; r1\tC3，+1.0V"), ["r1", "R2", "r1", "C3", "+1.0V"])


if __name__ == "__main__":
    unittest.main()
