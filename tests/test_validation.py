from dataclasses import asdict
import hashlib
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook

from bom_post_process.errors import BomError
from bom_post_process.validation import validate_workbook


class ValidationTests(unittest.TestCase):
    def setUp(self):
        self.temporary = TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.source = Path(self.temporary.name) / "input.xlsx"

    def write_rows(self, rows):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "AD export"
        sheet.append(["Comment", "Description", "Designator", "Footprint", "LibRef", "Quantity"])
        for row in rows:
            sheet.append(row)
        workbook.save(self.source)
        workbook.close()

    def test_supported_aliases_and_nc_suffixes(self):
        rows = []
        for index, footprint in enumerate(["C 0402L", "C0402", "CAP_0402", "0402C", "0402"]):
            rows.append(["0402-X5R104K250/NC", "", f"C{index + 1}", footprint, "", 1])
        for index, footprint in enumerate(["RES-0402", "0402R", "R0402", "0402"]):
            rows.append([" 0402k1002(NC) ", "", f"R{index + 1}", footprint, "", 1])
        self.write_rows(rows)
        before = hashlib.sha256(self.source.read_bytes()).hexdigest()
        report = validate_workbook(self.source)
        self.assertTrue(report.ok)
        self.assertEqual((report.checked, report.skipped), (9, 0))
        self.assertEqual([issue.code for issue in report.issues],
                         ["same_comment_different_footprints"] * 2)
        self.assertEqual([issue.row for issue in report.issues], [2, 7])
        self.assertEqual((report.errors, report.warnings), (0, 2))
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), before)

    def test_package_and_component_type_conflicts(self):
        self.write_rows([
            ["0402-X5R104K250", "", "C1", "C0603", "", 1],
            ["0402K1002", "", "R1", "0402C", "", 1],
            ["0603-X7R105K250", "", "C2", "R0402", "", 1],
        ])
        report = validate_workbook(self.source)
        self.assertFalse(report.ok)
        self.assertEqual((report.checked, report.errors, report.warnings), (3, 4, 0))
        self.assertEqual(
            [(issue.row, issue.code) for issue in report.issues],
            [(2, "package_mismatch"), (3, "component_type_mismatch"),
             (4, "package_mismatch"), (4, "component_type_mismatch")],
        )

    def test_unknown_models_and_ambiguous_footprints_are_unverified(self):
        self.write_rows([
            ["IC_M1206", "", "U1", "M1220", "", 1],
            ["1206B104K101NT(FH)", "", "C1", "C0402", "Cap", 1],
            ["0402K1002", "", "R1", "R0402/C0603", "", 1],
            ["0402K1002", "", "R2", "M0402", "", 1],
        ])
        report = validate_workbook(self.source)
        self.assertTrue(report.ok)
        self.assertEqual((report.checked, report.skipped, report.warnings), (0, 4, 4))
        self.assertEqual([issue.code for issue in report.issues],
                         ["unsupported_comment", "unknown_footprint", "unknown_footprint",
                          "same_comment_different_footprints"])
        self.assertEqual(report.issues[0].row, 3)

    def test_multiple_footprints_produce_one_group_warning_with_every_row(self):
        comment = "0402-X5R104K250"
        rows = [
            [comment, "", "C1,C2", "C0402", "Cap", 2],
            [comment, "", "C3", "C 0402L", "Cap", 1],
            [comment, "", "C4", "CAP_0402", "Cap", 1],
            [comment, "", "C5", "C0402", "Cap", 1],
        ]
        self.write_rows(rows)
        before = self.source.read_bytes()
        report = validate_workbook(self.source)
        self.assertEqual((report.checked, report.skipped, report.errors, report.warnings), (4, 0, 0, 1))
        self.assertTrue(report.ok)
        issue, = report.issues
        self.assertEqual((issue.severity, issue.code, issue.row),
                         ("warning", "same_comment_different_footprints", 2))
        self.assertEqual((issue.comment, issue.designator, issue.footprint), (comment, "C1,C2", "C0402"))
        self.assertIn(comment, issue.message)
        self.assertIn("请核对封装并合并对应 BOM 条目", issue.message)
        for row_number, row in enumerate(rows, start=2):
            self.assertRegex(issue.message, rf"第\s*{row_number}\s*行")
            self.assertIn(row[2], issue.message)
            self.assertIn(row[3], issue.message)
        self.assertEqual(self.source.read_bytes(), before)
        self.assertEqual(list(self.source.parent.iterdir()), [self.source])

    def test_case_and_outer_whitespace_are_ignored_for_group_comparison(self):
        self.write_rows([
            [" 0402k1002 ", "", "R1", " R0402 ", "", 1],
            ["0402K1002", "", "R2", "r0402", "", 1],
        ])
        report = validate_workbook(self.source)
        self.assertEqual((report.checked, report.skipped, report.errors, report.warnings), (2, 0, 0, 0))
        self.assertEqual(report.issues, ())

    def test_internal_footprint_whitespace_stays_significant_and_raw_text_is_preserved(self):
        self.write_rows([
            [" 0402k1002 ", "", "R1", " R0402 ", "", 1],
            ["0402K1002", "", "R2", "r0402", "", 1],
            ["0402k1002", "", "R3", "R 0402", "", 1],
        ])
        report = validate_workbook(self.source)
        issue, = report.issues
        self.assertEqual(issue.code, "same_comment_different_footprints")
        self.assertEqual((issue.comment, issue.footprint), (" 0402k1002 ", " R0402 "))
        self.assertIn(" 0402k1002 ", issue.message)
        for footprint in (" R0402 ", "r0402", "R 0402"):
            self.assertIn(footprint, issue.message)
        self.assertEqual((report.checked, report.skipped), (3, 0))

    def test_nc_suffixes_remain_separate_comments_including_parenthesized_form(self):
        self.write_rows([
            ["0402K1002", "", "R1", "R0402", "Res", 1],
            ["0402K1002/NC", "", "R2", "R 0402", "Res", 1],
            ["0402K1002(NC)", "", "R3", "RES_0402", "Res", 1],
        ])
        self.assertEqual(validate_workbook(self.source).issues, ())
        book = load_workbook(self.source)
        book.active.append([" 0402k1002/nc ", "", "R4", "RES-0402", "Res", 1])
        book.save(self.source)
        book.close()
        report = validate_workbook(self.source)
        issue, = report.issues
        self.assertEqual((issue.row, issue.comment), (3, "0402K1002/NC"))
        self.assertIn("R2", issue.message)
        self.assertIn("R4", issue.message)
        self.assertNotIn("R1", issue.message)
        self.assertNotIn("R3", issue.message)
        self.assertEqual((report.checked, report.skipped, report.errors, report.warnings), (4, 0, 0, 1))

    def test_manufacturer_models_and_polarized_capacitors_are_grouped_by_libref(self):
        self.write_rows([
            ["1206B104K101NT(FH)", "", "C1", "C1206", " Cap ", 1],
            ["1206B104K101NT(FH)", "", "C2", "CAP_1206", "CAPACITOR", 1],
            ["ERJ-3EKF1002V", "", "R1", "R0603", "Res1", 1],
            ["ERJ-3EKF1002V", "", "R2", "RES_0603", " resistor ", 1],
            ["EEE-1HA100WR", "", "C3", "CAP-A", " Cap Pol1 ", 1],
            ["EEE-1HA100WR", "", "C4", "CAP-B", "CAP POL2", 1],
        ])
        report = validate_workbook(self.source)
        self.assertEqual((report.checked, report.skipped, report.errors, report.warnings), (0, 6, 0, 9))
        self.assertEqual([issue.code for issue in report.issues],
                         ["unsupported_comment"] * 6 + ["same_comment_different_footprints"] * 3)
        self.assertEqual([issue.row for issue in report.issues[-3:]], [2, 4, 6])
        self.assertTrue(report.ok)

    def test_non_rc_rows_are_excluded_even_when_the_comment_matches_an_rc_row(self):
        self.write_rows([
            ["ChipX", "", "U1", "QFN32", "IC", 1],
            ["ChipX", "", "U2", "QFN48", "IC", 1],
            ["LED-A", "", "D1", "LED0402", "LED", 1],
            ["LED-A", "", "D2", "LED0603", "LED", 1],
            ["1206B104K101NT(FH)", "", "C1", "C1206", "", 1],
            ["1206B104K101NT(FH)", "", "C2", "C0402", "", 1],
            ["ManufacturerPart", "", "C3", "C0603", "Cap", 1],
            ["ManufacturerPart", "", "U3", "QFN48", "IC", 1],
        ])
        report = validate_workbook(self.source)
        self.assertEqual([issue.code for issue in report.issues], ["unsupported_comment"])
        self.assertEqual(report.issues[0].row, 8)
        self.assertEqual((report.checked, report.skipped, report.warnings), (0, 8, 1))

    def test_non_rc_components_skip_standard_comment_and_footprint_rules(self):
        self.write_rows([
            ["IC_M1206", "", "U1", "M1220", "IC", 1],
            ["LED-A", "", "D1", "LED0603", "LED", 1],
            ["CONN-10", "", "J1", "HEADER-10", "Connector", 1],
            ["MARK", "", "M1", "MARK-1MM", "MarkPoint", 1],
            ["MillMax-2501-2", "", "TP1", "TESTPOINT", "TestPoint", 1],
            ["VendorUnknown", "", "U2", "QFN", "", 1],
        ])
        before = self.source.read_bytes()
        report = validate_workbook(self.source)
        self.assertEqual(report.issues, ())
        self.assertEqual((report.checked, report.skipped, report.errors, report.warnings), (0, 6, 0, 0))
        self.assertTrue(report.ok)
        self.assertEqual(self.source.read_bytes(), before)

    def test_non_rc_components_still_receive_designator_quantity_and_missing_field_checks(self):
        self.write_rows([
            ["ChipX", "", "U1,U1", "QFN", "IC", 2],
            ["LED-A", "", "u1", "LED0603", "LED", 2],
            ["CONN-10", "", None, None, "Connector", 0],
            [None, "", "U2", "QFN", "IC", 1],
        ])
        report = validate_workbook(self.source)
        self.assertEqual([issue.code for issue in report.issues],
                         ["duplicate_designator", "duplicate_designator", "quantity_mismatch",
                          "missing_designator", "invalid_quantity", "missing_footprint", "missing_comment"])
        self.assertEqual((report.checked, report.skipped, report.errors, report.warnings), (0, 4, 7, 0))
        self.assertFalse(report.ok)

    def test_unknown_rc_and_polarized_capacitor_models_still_warn(self):
        self.write_rows([
            ["1206B104K101NT(FH)", "", "C1", "C1206", " Cap ", 1],
            ["ERJ-3EKF1002V", "", "R1", "R0603", " resistor ", 1],
            ["EEE-1HA100WR", "", "C2", "CAP-A", " Cap Pol1 ", 1],
        ])
        report = validate_workbook(self.source)
        self.assertEqual([issue.code for issue in report.issues], ["unsupported_comment"] * 3)
        self.assertEqual([issue.row for issue in report.issues], [2, 3, 4])
        self.assertTrue(all(issue.severity == "warning" for issue in report.issues))
        self.assertEqual((report.checked, report.skipped, report.errors, report.warnings), (0, 3, 0, 3))
        self.assertTrue(report.ok)

    def test_empty_comments_and_footprints_do_not_create_groups(self):
        self.write_rows([
            [None, "", "C1", "C0402", "Cap", 1],
            ["", "", "C2", "C0603", "Cap", 1],
            ["  ", "", "C3", "C0805", "Cap", 1],
            ["0402K1002", "", "R1", None, "Res", 1],
            ["0402K1002", "", "R2", "  ", "Res", 1],
            ["0402K1002", "", "R3", "R0402", "Res", 1],
        ])
        report = validate_workbook(self.source)
        self.assertEqual([issue.code for issue in report.issues],
                         ["missing_comment"] * 3 + ["missing_footprint"] * 2)
        self.assertEqual((report.checked, report.skipped, report.errors, report.warnings), (1, 5, 5, 0))

    def test_group_warning_is_appended_after_existing_row_errors(self):
        self.write_rows([
            ["0402K1002", "", "R1", "R0402", "Res", 1],
            ["0402K1002", "", "R1", "C0603", "Res", 2],
        ])
        report = validate_workbook(self.source)
        self.assertEqual([issue.code for issue in report.issues],
                         ["duplicate_designator", "quantity_mismatch", "package_mismatch",
                          "component_type_mismatch", "same_comment_different_footprints"])
        self.assertEqual((report.checked, report.skipped, report.errors, report.warnings), (2, 0, 4, 1))
        self.assertFalse(report.ok)
        self.assertEqual(report.issues[-1].row, 2)
        self.assertEqual(report.issues[-1].severity, "warning")

    def test_group_checks_only_the_selected_worksheet_without_modifying_workbook(self):
        self.write_rows([
            ["0402K1002", "", "R1", "R0402", "Res", 1],
            ["0402K1002", "", "R2", "R0402", "Res", 1],
        ])
        book = load_workbook(self.source)
        other = book.copy_worksheet(book.active)
        other.title = "另一张 AD"
        other["D3"] = "R 0402"
        book.save(self.source)
        book.close()
        before = self.source.read_bytes()
        report = validate_workbook(self.source, sheet="AD export")
        self.assertEqual(report.issues, ())
        report = validate_workbook(self.source, sheet="另一张 AD")
        self.assertEqual(report.sheet, "另一张 AD")
        self.assertEqual([issue.code for issue in report.issues], ["same_comment_different_footprints"])
        self.assertEqual((report.checked, report.skipped), (2, 0))
        self.assertEqual(self.source.read_bytes(), before)

    def test_duplicate_designators_and_quantity_mismatch(self):
        self.write_rows([
            ["0402K1002", "", "R1, R1", "R0402", "", 2],
            ["0402K1002", "", "r1; R2", "R0402", "", 3],
            ["0402K1002", "", "R3", "R0402", "", 0],
            ["0402K1002", "", "", "R0402", "", 1],
        ])
        report = validate_workbook(self.source)
        self.assertEqual([issue.code for issue in report.issues],
                         ["duplicate_designator", "duplicate_designator", "quantity_mismatch",
                          "invalid_quantity", "missing_designator"])
        self.assertIn("第 2 行", report.issues[1].message)

    def test_positive_integer_quantities_and_unusual_designators(self):
        self.write_rows([
            ["0402K1002", "", "+1.0V, VIN-1", "R0402", "", "2.0"],
            ["0402K1002", "", "R2，R3、R4\nR5", "R0402", "", 4.0],
            ["0402K1002", "", "R6 R7\tR8", "R0402", "", 3],
        ])
        report = validate_workbook(self.source)
        self.assertEqual(report.issues, ())

    def test_missing_values_are_errors(self):
        self.write_rows([
            [None, "", "C1", "C0402", "", 1],
            ["0402K1002", "", "R1", None, "", 1],
            ["0402K1002", "", "R2", "R0402", "", "NaN"],
            ["0402K1002", "", "R3", "R0402", "", 1.5],
        ])
        report = validate_workbook(self.source)
        self.assertFalse(report.ok)
        self.assertEqual([issue.code for issue in report.issues],
                         ["missing_comment", "missing_footprint", "invalid_quantity", "invalid_quantity"])

    def test_report_is_json_serializable_with_statistics(self):
        self.write_rows([["0402K1002", "", "R1", "R0603", "", 1]])
        report = validate_workbook(self.source, sheet="AD export")
        result = json.loads(json.dumps(asdict(report), ensure_ascii=False))
        self.assertEqual(result["sheet"], "AD export")
        self.assertEqual(result["errors"], 1)
        self.assertEqual(result["warnings"], 0)
        self.assertFalse(result["ok"])
        self.assertEqual(result["issues"][0]["designator"], "R1")

    def test_empty_table_is_not_reported_as_valid(self):
        self.write_rows([])
        with self.assertRaisesRegex(BomError, "没有 BOM 数据行"):
            validate_workbook(self.source)

    def test_supplied_workbook(self):
        source = Path(__file__).resolve().parents[1] / "input" / "DAQ_1010A_V1R0.xlsx"
        before = hashlib.sha256(source.read_bytes()).hexdigest()
        report = validate_workbook(source)
        self.assertEqual((report.checked, report.skipped, report.errors, report.warnings),
                         (91, 51, 0, 11))
        self.assertEqual({issue.code for issue in report.issues},
                         {"unsupported_comment", "same_comment_different_footprints"})
        self.assertEqual([issue.code for issue in report.issues[:5]], ["unsupported_comment"] * 5)
        self.assertEqual([issue.row for issue in report.issues[:5]], [11, 12, 13, 29, 43])
        self.assertEqual([issue.code for issue in report.issues[5:]], ["same_comment_different_footprints"] * 6)
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)


if __name__ == "__main__":
    unittest.main()
