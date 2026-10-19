from contextlib import redirect_stderr, redirect_stdout
from io import StringIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from bom_post_process.cli import main


class CliTests(unittest.TestCase):
    def setUp(self):
        self.temporary = TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)
        self.source = self.folder / "input.xlsx"
        workbook = Workbook()
        workbook.active.append(["Comment", "Description", "Designator", "Footprint", "LibRef", "Quantity"])
        workbook.active.append(["0402K1002", "", "R1", "R0402", "Res", 1])
        workbook.save(self.source)
        workbook.close()

    def run_cli(self, *args):
        stdout, stderr = StringIO(), StringIO()
        with redirect_stdout(stdout), redirect_stderr(stderr):
            code = main([str(arg) for arg in args])
        return code, stdout.getvalue(), stderr.getvalue()

    def test_validate_reports_conflict_and_preserves_input(self):
        workbook = load_workbook(self.source)
        workbook.active["D2"] = "R0603"
        workbook.save(self.source)
        workbook.close()
        before = self.source.read_bytes()
        report = self.folder / "report.json"
        code, _, stderr = self.run_cli("validate", self.source, "-o", report)
        self.assertEqual(code, 1)
        self.assertIn("0603", stderr)
        data = json.loads(report.read_text(encoding="utf-8"))
        self.assertFalse(data["ok"])
        self.assertEqual(data["errors"], 1)
        self.assertEqual(self.source.read_bytes(), before)

    def test_merge_warnings_are_not_truncated_and_json_contains_all_details(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Comment", "Description", "Designator", "Footprint", "LibRef", "Quantity"])
        for index in range(10):
            sheet.append([f"VendorCap-{index}", "", f"C{index}", "C0603", "Cap", 1])
        for index in range(9):
            model = f"0402K{1000 + index}"
            sheet.append([model, "", f"R{index + 1}A", "R0402", "Res", 1])
            sheet.append([model, "", f"R{index + 1}B", "0402R", "Res", 1])
        workbook.save(self.source)
        workbook.close()
        before = self.source.read_bytes()
        output = self.folder / "report.json"
        code, stdout, stderr = self.run_cli("validate", self.source, "-o", output)
        self.assertEqual(code, 0)
        self.assertIn("错误 0，提示 19", stdout)
        self.assertEqual(stderr.count("请核对封装并合并对应 BOM 条目"), 9)
        self.assertEqual(stderr.count("未验证封装一致性"), 8)
        self.assertIn("另有 2 条提示", stderr)
        data = json.loads(output.read_text(encoding="utf-8"))
        self.assertEqual((data["checked"], data["skipped"], data["warnings"]), (18, 10, 19))
        self.assertTrue(data["ok"])
        self.assertEqual(len(data["issues"]), 19)
        groups = [item for item in data["issues"] if item["code"] == "same_comment_different_footprints"]
        self.assertEqual(len(groups), 9)
        for index, group in enumerate(groups):
            self.assertEqual(group["row"], 12 + 2 * index)
            self.assertEqual(group["severity"], "warning")
            self.assertEqual(group["comment"], f"0402K{1000 + index}")
            self.assertEqual(group["designator"], f"R{index + 1}A")
            self.assertEqual(group["footprint"], "R0402")
            self.assertIn(group["message"], stderr)
            for value in (f"第 {12 + 2 * index} 行", f"第 {13 + 2 * index} 行",
                          f"R{index + 1}A", f"R{index + 1}B", "R0402", "0402R"):
                self.assertIn(value, group["message"])
        self.assertEqual(self.source.read_bytes(), before)

    def test_non_rc_components_do_not_emit_comment_warnings_in_console_or_json(self):
        workbook = load_workbook(self.source)
        for row in [
            ["ChipX", "", "U1", "QFN32", "IC", 1],
            ["ConnectorX", "", "J1", "HDR2", "Connector", 1],
            ["LED-RED", "", "D1", "LED0603", "LED", 1],
            ["Testpoint", "", "TP1", "TESTPOINT", "Testpoint", 1],
            ["MarkX", "", "M1", "MARK", "Markpoint", 1],
            ["VendorCap", "", "C1", "C0603", "Cap", 1],
        ]:
            workbook.active.append(row)
        workbook.save(self.source)
        workbook.close()
        before = self.source.read_bytes()
        output = self.folder / "report.json"
        code, stdout, stderr = self.run_cli("validate", self.source, "-o", output)
        self.assertEqual(code, 0)
        self.assertIn("错误 0，提示 1", stdout)
        self.assertEqual(stderr.count("未验证封装一致性"), 1)
        self.assertIn("[C1]", stderr)
        for designator in ("U1", "J1", "D1", "TP1", "M1"):
            self.assertNotIn(f"[{designator}]", stderr)
        data = json.loads(output.read_text(encoding="utf-8"))
        self.assertEqual((data["checked"], data["skipped"], data["warnings"]), (1, 6, 1))
        self.assertTrue(data["ok"])
        self.assertEqual([(issue["code"], issue["designator"]) for issue in data["issues"]],
                         [("unsupported_comment", "C1")])
        self.assertEqual(self.source.read_bytes(), before)

    def test_merge_warning_does_not_prevent_exports_or_merge_rows(self):
        workbook = load_workbook(self.source)
        workbook.active.append(["0402K1002", "", "R2", "0402R", "Res", 1])
        workbook.save(self.source)
        workbook.close()
        before = self.source.read_bytes()
        code, _, stderr = self.run_cli("validate", self.source)
        self.assertEqual(code, 0)
        self.assertIn("请核对封装并合并对应 BOM 条目", stderr)
        parsed = self.folder / "parsed.xlsx"
        soldering = self.folder / "soldering.xlsx"
        bom = self.folder / "BOM.xlsx"
        self.assertEqual(self.run_cli("gen_value", self.source, "-o", parsed)[0], 0)
        self.assertEqual(self.run_cli("gen_soldering", parsed, "-o", soldering)[0], 0)
        self.assertEqual(self.run_cli("gen_bom", parsed, "-o", bom)[0], 0)
        self.assertEqual(self.source.read_bytes(), before)
        for path, first_material_row, designator_column in ((parsed, 2, 3), (soldering, 2, 5), (bom, 3, 5)):
            book = load_workbook(path)
            try:
                self.assertEqual([book.active.cell(row, designator_column).value
                                  for row in range(first_material_row, book.active.max_row + 1)],
                                 ["R1", "R2"])
            finally:
                book.close()

    def test_package_error_still_fails_when_merge_warning_is_present(self):
        workbook = load_workbook(self.source)
        workbook.active.append(["0402K1002", "", "R2", "R0603", "Res", 1])
        workbook.save(self.source)
        workbook.close()
        code, stdout, stderr = self.run_cli("validate", self.source)
        self.assertEqual(code, 1)
        self.assertIn("错误 1，提示 1", stdout)
        self.assertIn("Footprint 指定 0603", stderr)
        self.assertIn("请核对封装并合并对应 BOM 条目", stderr)

    def test_gen_value_and_collision(self):
        output = self.folder / "nested" / "parsed.xlsx"
        code, _, _ = self.run_cli("gen_value", self.source, "-o", output)
        self.assertEqual(code, 0)
        workbook = load_workbook(output)
        self.assertEqual([cell.value for cell in workbook.active[1]][-4:],
                         ["name", "note", "not_bought", "not_welded"])
        self.assertEqual(workbook.active["G2"].value, "贴片电阻")
        self.assertEqual(workbook.active["H2"].value, "10k/0402/1%")
        self.assertEqual(len(workbook.worksheets), 1)
        self.assertNotIn("焊接表", workbook.sheetnames)
        workbook.close()
        before = output.read_bytes()
        code, _, stderr = self.run_cli("gen_value", self.source, "-o", output)
        self.assertEqual(code, 1)
        self.assertIn("已存在", stderr)
        self.assertEqual(output.read_bytes(), before)

    def test_all_is_not_a_command(self):
        with self.assertRaises(SystemExit) as caught:
            self.run_cli("all", self.source)
        self.assertEqual(caught.exception.code, 2)

    def test_manual_selection_survives_soldering_export(self):
        parsed = self.folder / "parsed.xlsx"
        soldering = self.folder / "soldering.xlsx"
        code, _, _ = self.run_cli("gen_value", self.source, "-o", parsed)
        self.assertEqual(code, 0)
        workbook = load_workbook(parsed)
        note = "10k/0402/0.1%\n选型：低温漂，指定厂家 & 型号"
        workbook.active["G2"] = "精密电阻"
        workbook.active["H2"] = note
        workbook.save(parsed)
        workbook.close()
        code, _, _ = self.run_cli("gen_soldering", parsed, "-o", soldering)
        self.assertEqual(code, 0)
        workbook = load_workbook(soldering)
        self.assertEqual(workbook["焊接表"]["C2"].value, "精密电阻")
        self.assertEqual(workbook["焊接表"]["F2"].value, note)
        workbook.close()

    def test_reviewed_file_branches_into_independent_soldering_and_bom_exports(self):
        workbook = load_workbook(self.source)
        for index in range(2, 5):
            workbook.active.append(["0402K1002", "", f"R{index}", "R0402", "Res", 1])
        workbook.save(self.source)
        workbook.close()
        parsed = self.folder / "parsed.xlsx"
        soldering = self.folder / "soldering.xlsx"
        bom = self.folder / "BOM.xlsx"
        self.assertEqual(self.run_cli("gen_value", self.source, "-o", parsed)[0], 0)
        workbook = load_workbook(parsed)
        sheet = workbook.active
        sheet["I3"], sheet["J4"] = "仓库已有，不购买", "暂不焊接"
        sheet["I5"], sheet["J5"] = "无需采购", "无需焊接"
        sheet["G4"], sheet["H4"] = "人工名称", "采购时保留本备注\n第二行"
        workbook.save(parsed)
        workbook.close()
        before = parsed.read_bytes()
        self.assertEqual(self.run_cli("gen_soldering", parsed, "-o", soldering)[0], 0)
        self.assertEqual(self.run_cli("gen_bom", parsed, "-o", bom)[0], 0)
        with_soldering = load_workbook(soldering)
        with_bom = load_workbook(bom)
        self.addCleanup(with_soldering.close)
        self.addCleanup(with_bom.close)
        self.assertEqual(with_soldering.sheetnames, ["焊接表"])
        self.assertEqual(with_bom.sheetnames, ["BOM"])
        self.assertEqual([row[4] for row in with_soldering.active.iter_rows(min_row=2, values_only=True)],
                         ["R1", "R2"])
        records = list(with_bom.active.iter_rows(min_row=2, values_only=True))
        self.assertEqual(records[0][1:4], ("PCB制板", "PCB制板", None))
        self.assertEqual([row[4] for row in records[1:]], ["R1", "R3"])
        self.assertEqual(records[2][2], "人工名称")
        self.assertEqual(records[2][5], "采购时保留本备注\n第二行")
        self.assertEqual(parsed.read_bytes(), before)

    def test_nc_prefill_can_be_manually_cleared_before_soldering(self):
        workbook = load_workbook(self.source)
        workbook.active["A2"] = "0402K1002/NC"
        workbook.save(self.source)
        workbook.close()
        parsed = self.folder / "parsed.xlsx"
        self.assertEqual(self.run_cli("gen_value", self.source, "-o", parsed)[0], 0)
        workbook = load_workbook(parsed)
        self.assertEqual(workbook.active["J2"].value, "NC")
        self.assertIn(workbook.active["I2"].value, (None, ""))
        workbook.active["J2"] = None
        workbook.save(parsed)
        workbook.close()
        soldering = self.folder / "soldering.xlsx"
        self.assertEqual(self.run_cli("gen_soldering", parsed, "-o", soldering)[0], 0)
        workbook = load_workbook(soldering)
        self.assertEqual(workbook.active["E2"].value, "R1")
        workbook.close()

    def test_standard_export_is_an_independent_command(self):
        from bom_post_process.standard import StandardReport

        word = self.folder / "result.docx"
        report = StandardReport(word, word.with_suffix(".pdf"), word.with_suffix(".json"), 1, 1, "BOM")
        with patch("bom_post_process.standard.export_standard_bom", return_value=report) as export:
            code, _, _ = self.run_cli("gen_bom_standard", self.source, "-o", word)
        self.assertEqual(code, 0)
        export.assert_called_once_with(self.source, word, sheet=None, template=None, config=None)

    def test_missing_file_is_a_friendly_error(self):
        code, _, stderr = self.run_cli("validate", self.folder / "missing.xlsx")
        self.assertEqual(code, 1)
        self.assertIn("不存在", stderr)
        self.assertNotIn("Traceback", stderr)


if __name__ == "__main__":
    unittest.main()
