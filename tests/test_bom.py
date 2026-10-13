from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from bom_post_process.bom import gen_bom
from bom_post_process.errors import BomError
from bom_post_process.standard import RESOURCES, Template, read_bom, read_exported_content
from bom_post_process.values import gen_value
from bom_post_process.workbook import AD_HEADERS, BOM_HEADERS


class BomTests(unittest.TestCase):
    def setUp(self):
        self.temporary = TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)
        self.source = self.folder / "样例_已解析.xlsx"
        self.output = self.folder / "BOM.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "已确认"
        sheet.append([*AD_HEADERS, "name", "note", "not_bought", "not_welded"])
        for row in [
            ["ChipX", "描述不导出", "U1", "QFN", "IC", 1,
             "手工芯片名称", "  手工选型：指定厂家 & 型号\n工作温度 -40~85℃  ", None, "不用焊但要买"],
            ["0402K1003/NC", "", "R1", "R0402", "Res", 1,
             "贴片电阻", "备件", None, "NC"],
            ["MARK", "", "M1", "MARK", "MarkPoint", 1,
             "定位点", "板上图形", "不买", None],
        ]:
            sheet.append(row)
        sheet["H2"].font = Font(name="宋体", size=12, bold=True)
        sheet.column_dimensions["H"].width = 48
        book.save(self.source)
        book.close()

    def edit(self, action, destination=None):
        book = load_workbook(self.source)
        action(book)
        path = destination or self.source
        book.save(path)
        book.close()
        return path

    def test_purchase_flag_only_with_blank_quantity_pcb_placeholder(self):
        before = self.source.read_bytes()
        report = gen_bom(self.source, self.output)
        self.assertEqual(self.source.read_bytes(), before)
        self.assertEqual((report.written, report.excluded, report.pcb_placeholder), (3, 1, 1))
        book = load_workbook(self.output)
        try:
            self.assertEqual(book.sheetnames, ["BOM"])
            values = list(book.active.values)
            self.assertEqual(values[0], BOM_HEADERS)
            self.assertEqual(values[1], (1, "PCB制板", "PCB制板", None, None, "PCB制板的占位符"))
            self.assertEqual(values[2][0:5], (2, "ChipX", "手工芯片名称", 1, "U1"))
            self.assertEqual(values[3][0:5], (3, "0402K1003/NC", "贴片电阻", 1, "R1"))
            self.assertTrue(book.active["F3"].font.bold)
            self.assertEqual(book.active.column_dimensions["F"].width, 48)
        finally:
            book.close()

    def test_existing_pcb_is_not_duplicated_and_keeps_position_and_user_values(self):
        for index, field in enumerate(("comment", "name")):
            with self.subTest(field=field):
                def change(book):
                    row = ["用户板号", "", "PCB1", "", "", 7, "定制电路板", "用户板厚1.6mm", None, None]
                    row[0 if field == "comment" else 6] = "PCB制板"
                    book.active.append(row)
                source = self.edit(change, self.folder / f"pcb_{index}.xlsx")
                output = self.folder / f"pcb_out{index}.xlsx"
                report = gen_bom(source, output)
                self.assertEqual((report.written, report.pcb_placeholder), (3, 0))
                book = load_workbook(output)
                try:
                    self.assertEqual(book.active["B2"].value, "ChipX")
                    self.assertEqual(book.active["D4"].value, 7)
                    self.assertEqual(book.active["F4"].value, "用户板厚1.6mm")
                finally:
                    book.close()

    def test_explicitly_excluded_pcb_is_never_reinserted(self):
        self.edit(lambda book: book.active.append(
            ["PCB制板", "", "", "", "", 3, "PCB制板", "已有板无需采购", "不买", None]))
        report = gen_bom(self.source, self.output)
        self.assertEqual((report.written, report.excluded, report.pcb_placeholder), (2, 2, 0))
        book = load_workbook(self.output)
        try:
            self.assertFalse(any(row[1] == "PCB制板" for row in list(book.active.values)[1:]))
        finally:
            book.close()

    def test_zero_false_and_nonblank_text_purchase_markers_exclude(self):
        for index, (value, excluded) in enumerate([(None, 1), (" \t", 1), (0, 2), (False, 2), ("False", 2)]):
            with self.subTest(value=value):
                source = self.edit(lambda book: setattr(book.active["I2"], "value", value),
                                   self.folder / f"marker{index}.xlsx")
                report = gen_bom(source, self.folder / f"marker_out{index}.xlsx")
                self.assertEqual(report.excluded, excluded)

    def test_soldering_formula_or_error_does_not_affect_purchase_export(self):
        for index, value in enumerate(("=1+1", "#REF!")):
            source = self.edit(lambda book: setattr(book.active["J2"], "value", value),
                               self.folder / f"ignored{index}.xlsx")
            report = gen_bom(source, self.folder / f"ignored_out{index}.xlsx")
            self.assertEqual(report.written, 3)

    def test_formula_in_purchase_marker_rejected_without_output(self):
        self.edit(lambda book: setattr(book.active["I2"], "value", "=1+1"))
        with self.assertRaisesRegex(BomError, "not_bought.*公式"):
            gen_bom(self.source, self.output)
        self.assertFalse(self.output.exists())

    def test_reordered_columns_and_manual_text_carry_through_to_word(self):
        note = "  指定选型 & 厂商\n耐温 -40~85℃  "
        book = Workbook()
        sheet = book.active
        sheet.append(["note", "Maker", "Quantity", "LibRef", "Comment", "Designator", "Description",
                      "Footprint", "not_welded", "name", "not_bought"])
        sheet.append([note, "厂家A", 1, "IC", "ChipX", "U1", "原描述", "QFN", "不焊接", "手工名称", None])
        book.save(self.source)
        book.close()
        gen_bom(self.source, self.output)
        _, records = read_bom(self.output)
        self.assertEqual(records, [["1", "PCB制板", "PCB制板", "", "", "PCB制板的占位符"],
                                   ["2", "ChipX", "手工名称", "1", "U1", note]])
        template = Template(RESOURCES / "bom_template.docx")
        word = self.folder / "reviewed.docx"
        template.write(word, [records])
        material, categories = read_exported_content(word, template)
        self.assertEqual(material, records)
        self.assertEqual(categories, [(0, "外协件"), (1, "外购件")])

    def test_default_name_and_output_never_overwritten(self):
        report = gen_bom(self.source)
        self.assertEqual(report.output.name, "样例_BOM.xlsx")
        for output in (self.source, report.output):
            before = output.read_bytes()
            with self.assertRaises(BomError):
                gen_bom(self.source, output)
            self.assertEqual(output.read_bytes(), before)

    def test_real_parsed_input_includes_nc_purchases_and_one_pcb_placeholder(self):
        source = Path(__file__).resolve().parents[1] / "input" / "DAQ_1010A_V1R0.xlsx"
        parsed = self.folder / "parsed.xlsx"
        gen_value(source, parsed)
        report = gen_bom(parsed, self.output)
        self.assertEqual((report.written, report.excluded, report.pcb_placeholder), (143, 0, 1))


if __name__ == "__main__":
    unittest.main()
