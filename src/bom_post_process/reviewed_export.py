"""Map manually reviewed AD rows to the six-column downstream tables."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from .errors import BomError
from .workbook import (
    AD_HEADERS, BOM_HEADERS, cell_text, copy_style, copy_table_layout, copy_table_row,
    open_workbook, prepare_output, save_workbook_exclusive, select_reviewed_sheet, set_text,
)


@dataclass(frozen=True)
class ExportReport:
    output: Path
    written: int
    excluded: int
    source_sheet: str
    sheet: str
    pcb_placeholder: int = 0


def _check_value(cell, label: str) -> None:
    if cell.data_type in {"f", "e"}:
        raise BomError(f"{label} 单元格 {cell.coordinate} 包含公式或 Excel 错误，请先转为确认后的值")


def _copy_column_dimensions(source, target, columns) -> None:
    for output_column, input_column in enumerate(columns, start=1):
        copied = target.column_dimensions[get_column_letter(output_column)]
        copied.width = source.sheet_format.defaultColWidth or 13
        copied.hidden = False
        for dimension in source.column_dimensions.values():
            first = dimension.min or column_index_from_string(dimension.index)
            last = dimension.max or first
            if first <= input_column <= last:
                for attribute in ("width", "hidden", "bestFit", "outlineLevel", "collapsed"):
                    setattr(copied, attribute, getattr(dimension, attribute))
                copy_style(dimension, copied)
                break


def export_reviewed(
    input_path: Path | str, output_path: Path | str | None, sheet: str | None,
    *, flag: str, title: str, add_pcb: bool = False,
) -> ExportReport:
    source_path = Path(input_path).expanduser().resolve()
    default = source_path.with_name(source_path.stem.removesuffix("_已解析") + f"_{title}.xlsx")
    output = prepare_output(source_path, output_path or default, ".xlsx")
    source_book = open_workbook(source_path)
    result = Workbook()
    try:
        source, headers = select_reviewed_sheet(source_book, sheet)
        fields = ("comment", "comment", "name", "quantity", "designator", "note")
        columns = [headers[name] for name in fields]
        used_columns = set(columns) | {headers[flag]}
        if any(any(merged.min_col <= column <= merged.max_col for column in used_columns)
               for merged in source.merged_cells.ranges):
            raise BomError("待导出字段包含合并单元格，请取消合并后再导出")
        target = result.active
        target.title = title
        result.loaded_theme = source_book.loaded_theme
        result.epoch = source_book.epoch
        copy_table_layout(source, target)
        _copy_column_dimensions(source, target, columns)
        copy_table_row(source, target, 1, 1, columns)
        for column, label in enumerate(BOM_HEADERS, start=1):
            set_text(target.cell(1, column), label)

        # Flags alone do not create material rows. The other flag is never read.
        data_columns = [headers[name.casefold()] for name in (*AD_HEADERS, "name", "note")]
        rows = [row for row in range(2, source.max_row + 1)
                if any(source.cell(row, column).value not in (None, "") for column in data_columns)]
        if not rows:
            raise BomError("已解析的表格没有物料数据")
        # Check every reviewed row so an explicitly excluded PCB is never added back.
        has_pcb = any(cell_text(source.cell(row, headers[name]).value).strip() == "PCB制板"
                      for row in rows for name in ("comment", "name"))
        pcb_placeholder = int(add_pcb and not has_pcb)
        written = pcb_placeholder
        excluded = 0
        if pcb_placeholder:
            for output_column, source_column in enumerate(columns, start=1):
                copy_style(source.cell(rows[0], source_column), target.cell(2, output_column))
            for column, value in enumerate((1, "PCB制板", "PCB制板", None, "", "PCB制板的占位符"), start=1):
                target.cell(2, column, value)
        for row in rows:
            marker = source.cell(row, headers[flag])
            _check_value(marker, flag)
            if cell_text(marker.value).strip():
                excluded += 1
                continue
            for name, column in zip(fields, columns):
                _check_value(source.cell(row, column), name)
            written += 1
            copy_table_row(source, target, row, written + 1, columns)
            target.cell(written + 1, 1, written)
        target.auto_filter.ref = f"A1:F{target.max_row}"
        target.print_title_rows = "1:1"
        target.print_area = f"A1:F{target.max_row}"
        save_workbook_exclusive(result, output)
        return ExportReport(output, written, excluded, source.title, title, pcb_placeholder)
    finally:
        source_book.close()
        result.close()
