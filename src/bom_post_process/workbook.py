"""Shared workbook loading and header-based AD table discovery."""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import re
from xml.etree import ElementTree as ET
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from .errors import BomError


AD_HEADERS = ("Comment", "Description", "Designator", "Footprint", "LibRef", "Quantity")
BOM_HEADERS = ("序号", "规格型号/代号", "名称", "数量", "项目代号", "备注")


@dataclass(frozen=True)
class AdRow:
    row: int
    comment: str
    description: str
    designator: str
    footprint: str
    libref: str
    quantity: object


@dataclass(frozen=True)
class AdTable:
    sheet: str
    rows: tuple[AdRow, ...]


def cell_text(value: object) -> str:
    return "" if value is None else str(value)


def tokenize_designators(value: object) -> list[str]:
    """Split common AD separators, retaining duplicates and case for validation."""
    return [part for part in re.split(r"[,，;；、\s]+", cell_text(value).strip()) if part]


def split_designators(value: object) -> set[str]:
    return {part.upper() for part in tokenize_designators(value)}


def open_workbook(input_path: Path | str):
    source = Path(input_path).expanduser().resolve()
    if source.suffix.lower() != ".xlsx":
        raise BomError("输入文件必须为 .xlsx 格式")
    if not source.is_file():
        raise BomError(f"输入文件不存在：{source}")
    try:
        return load_workbook(BytesIO(source.read_bytes()), data_only=False)
    except (BadZipFile, InvalidFileException, ET.ParseError, KeyError, ValueError) as error:
        raise BomError(f"无法读取 Excel 文件：{error}") from error


def header_map(worksheet) -> dict[str, int]:
    result: dict[str, int] = {}
    for cell in worksheet[1]:
        value = cell_text(cell.value).strip().casefold()
        if value:
            if value in result:
                raise BomError(f"工作表“{worksheet.title}”存在重复表头：{cell.value}")
            result[value] = cell.column
    return result


def select_ad_sheet(workbook, sheet: str | None = None):
    required = {name.casefold() for name in AD_HEADERS}
    if sheet is not None:
        if sheet not in workbook.sheetnames:
            raise BomError(f"找不到工作表：{sheet}")
        selected = workbook[sheet]
        headers = header_map(selected)
        missing = [name for name in AD_HEADERS if name.casefold() not in headers]
        if missing:
            raise BomError(f"工作表“{sheet}”缺少 AD 表头：{', '.join(missing)}")
        return selected, headers
    candidates = []
    for candidate in workbook.worksheets:
        # Ignore unrelated layouts (including cover sheets with repeated labels).
        labels = {cell_text(cell.value).strip().casefold() for cell in candidate[1]}
        if required <= labels:
            candidates.append((candidate, header_map(candidate)))
    if not candidates:
        raise BomError("找不到 AD 原始表；需要表头：" + ", ".join(AD_HEADERS))
    if len(candidates) > 1:
        raise BomError("存在多个 AD 原始表，请使用 --sheet 指定工作表")
    return candidates[0]


def read_ad_table(workbook, sheet: str | None = None) -> AdTable:
    worksheet, headers = select_ad_sheet(workbook, sheet)
    indexes = [headers[name.casefold()] for name in AD_HEADERS]
    rows: list[AdRow] = []
    for row_number in range(2, worksheet.max_row + 1):
        cells = [worksheet.cell(row_number, column) for column in indexes]
        if all(cell.value is None or cell.value == "" for cell in cells):
            continue
        for cell in cells:
            if cell.data_type == "f":
                raise BomError(f"工作表“{worksheet.title}”{cell.coordinate} 包含公式，请先转为值")
            if cell.data_type == "e":
                raise BomError(f"工作表“{worksheet.title}”{cell.coordinate} 包含 Excel 错误：{cell.value}")
        values = [cell.value for cell in cells]
        rows.append(AdRow(row_number, *(cell_text(value) for value in values[:5]), values[5]))
    return AdTable(worksheet.title, tuple(rows))


def read_ad_workbook(input_path: Path | str, sheet: str | None = None) -> AdTable:
    workbook = open_workbook(input_path)
    try:
        return read_ad_table(workbook, sheet)
    finally:
        workbook.close()


def prepare_output(input_path: Path | str, output_path: Path | str, suffix: str) -> Path:
    source = Path(input_path).expanduser().resolve()
    output = Path(output_path).expanduser().resolve()
    if output.suffix.lower() != suffix:
        raise BomError(f"输出文件必须为 {suffix} 格式")
    if source == output:
        raise BomError("输出路径不能与输入路径相同")
    if output.exists():
        raise BomError(f"输出文件已存在，请指定其他路径：{output}")
    return output


def save_workbook_exclusive(workbook, output: Path) -> None:
    buffer = BytesIO()
    workbook.save(buffer)
    save_bytes_exclusive(buffer.getvalue(), output)


def save_bytes_exclusive(data: bytes, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("xb") as stream:
        try:
            stream.write(data)
            stream.flush()
        except OSError:
            stream.close()
            output.unlink(missing_ok=True)
            raise


def select_reviewed_sheet(workbook, sheet: str | None = None):
    """Find the parsed AD table with all four manually editable columns."""
    required = {name.casefold() for name in AD_HEADERS} | {"name", "note", "not_bought", "not_welded"}
    if sheet is not None and sheet not in workbook.sheetnames:
        raise BomError(f"找不到工作表：{sheet}")
    candidates = []
    for candidate in ([workbook[sheet]] if sheet else workbook.worksheets):
        labels = {cell_text(cell.value).strip().casefold() for cell in candidate[1]}
        if required <= labels:
            candidates.append((candidate, header_map(candidate)))
    if len(candidates) > 1:
        raise BomError("存在多个已解析的 AD 表，请使用 --sheet 指定已编辑的工作表")
    if not candidates:
        raise BomError("缺少已解析的 AD 表及 name、note、not_bought、not_welded 四列；"
                       "请先运行 gen_value 并人工修改后再导出")
    return candidates[0]


def copy_style(source, target) -> None:
    for attribute in ("font", "fill", "border", "alignment", "number_format", "protection"):
        setattr(target, attribute, copy(getattr(source, attribute)))
    target.quotePrefix = source.quotePrefix


def copy_table_layout(source, target) -> None:
    target.sheet_format = copy(source.sheet_format)
    target.sheet_properties = copy(source.sheet_properties)
    target.sheet_view.showGridLines = source.sheet_view.showGridLines
    target.freeze_panes = source.freeze_panes
    for dimension in source.column_dimensions.values():
        first = dimension.min or column_index_from_string(dimension.index)
        last = dimension.max or first
        for column in range(first, min(last, 6) + 1):
            copied = target.column_dimensions[get_column_letter(column)]
            for attribute in ("width", "hidden", "bestFit", "outlineLevel", "collapsed"):
                setattr(copied, attribute, getattr(dimension, attribute))
            copy_style(dimension, copied)
    target.page_margins = copy(source.page_margins)
    target.page_setup = copy(source.page_setup)
    target.page_setup.id = None
    target.print_options = copy(source.print_options)


def copy_table_row(source, target, source_row: int, output_row: int, columns=None) -> None:
    dimension = source.row_dimensions.get(source_row)
    if dimension is not None:
        copied = target.row_dimensions[output_row]
        for attribute in ("height", "hidden", "outlineLevel", "collapsed", "thickTop", "thickBot"):
            setattr(copied, attribute, getattr(dimension, attribute))
        copy_style(dimension, copied)
    for column, source_column in enumerate(columns or range(1, 7), start=1):
        cell = source.cell(source_row, source_column)
        copied = target.cell(output_row, column, cell.value)
        copied.data_type = cell.data_type
        copy_style(cell, copied)
        if cell.comment is not None:
            copied.comment = copy(cell.comment)
        if cell.hyperlink is not None:
            copied.hyperlink = copy(cell.hyperlink)


def set_text(cell, value: str) -> None:
    cell.value = value
    cell.data_type = "s"
