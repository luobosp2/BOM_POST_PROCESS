"""Patch review cells inside an XLSX, retaining other ZIP members verbatim."""

from __future__ import annotations

from io import BytesIO
import posixpath
from xml.dom import Node, minidom
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl.utils import get_column_letter, range_boundaries

from .errors import BomError


SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _worksheet_part(archive: ZipFile, title: str) -> str:
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    sheet = next((s for s in workbook.findall(f"{{{SHEET_NS}}}sheets/{{{SHEET_NS}}}sheet")
                  if s.get("name") == title), None)
    if sheet is None:
        raise BomError(f"找不到工作表：{title}")
    relation_id = sheet.get(f"{{{REL_NS}}}id")
    relations = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    for relation in relations:
        if relation.get("Id") == relation_id:
            target = relation.get("Target", "")
            return posixpath.normpath(target.lstrip("/") if target.startswith("/") else "xl/" + target)
    raise BomError(f"找不到工作表的数据文件：{title}")


def _children(node: Node, name: str) -> list:
    return [child for child in node.childNodes
            if child.nodeType == Node.ELEMENT_NODE
            and child.namespaceURI == SHEET_NS and child.localName == name]


def _patch_sheet(xml: bytes, column: int, values: dict[int, str],
                 reference_column: int, width: int | None) -> bytes:
    with minidom.parseString(xml) as document:
        root = document.documentElement
        prefix = f"{root.prefix}:" if root.prefix else ""

        def element(name: str):
            return document.createElementNS(SHEET_NS, prefix + name)

        data = _children(root, "sheetData")[0]
        rows = {int(row.getAttribute("r")): row for row in _children(data, "row")}
        letter = get_column_letter(column)
        reference_letter = get_column_letter(reference_column)
        for number, value in values.items():
            row = rows[number]
            cells = _children(row, "c")
            address = f"{letter}{number}"
            cell = next((c for c in cells if c.getAttribute("r") == address), None)
            if cell is None:
                cell = element("c")
                cell.setAttribute("r", address)
                donor = next((c for c in cells if c.getAttribute("r") == f"{reference_letter}{number}"), None)
                if donor is not None and donor.hasAttribute("s"):
                    cell.setAttribute("s", donor.getAttribute("s"))
                following = next((c for c in cells
                                  if range_boundaries(c.getAttribute("r"))[0] > column), None)
                row.insertBefore(cell, following if following is not None else (cells[-1].nextSibling if cells else row.firstChild))
            for child in list(cell.childNodes):
                if (child.nodeType == Node.ELEMENT_NODE and child.namespaceURI == SHEET_NS
                        and child.localName in {"f", "v", "is"}):
                    cell.removeChild(child)
                    child.unlink()
            cell.setAttribute("t", "inlineStr")
            inline = element("is")
            text = element("t")
            text.setAttribute("xml:space", "preserve")
            text.appendChild(document.createTextNode(value))
            inline.appendChild(text)
            cell.insertBefore(inline, cell.firstChild)
            if row.hasAttribute("spans"):
                bounds = [int(part) for span in row.getAttribute("spans").split() for part in span.split(":")]
                row.setAttribute("spans", f"{min(*bounds, column)}:{max(*bounds, column)}")

        if width is not None:
            dimensions = _children(root, "dimension")
            if dimensions:
                left, top, right, bottom = range_boundaries(dimensions[0].getAttribute("ref"))
                dimensions[0].setAttribute("ref", f"{get_column_letter(left)}{top}:{get_column_letter(max(right, column))}{bottom}")
            # Only size the added column. Existing widths, row heights and styles stay intact.
            existing_cols = _children(root, "cols")
            cols = existing_cols[0] if existing_cols else element("cols")
            if not existing_cols:
                root.insertBefore(cols, data)
            if not any(int(c.getAttribute("min")) <= column <= int(c.getAttribute("max")) for c in _children(cols, "col")):
                added = element("col")
                for key, value in {"min": str(column), "max": str(column), "width": str(width), "customWidth": "1"}.items():
                    added.setAttribute(key, value)
                following = next((c for c in _children(cols, "col") if int(c.getAttribute("min")) > column), None)
                cols.insertBefore(added, following)
        return document.toxml(encoding="UTF-8", standalone=document.standalone)


def _range_reaches_column(value: str, column: int) -> bool:
    for area in value.split():
        try:
            left, _, right, _ = range_boundaries(area.replace("$", ""))
        except ValueError:
            return True
        if left is None or right is None or right >= column:
            return True
    return False


def _check_legacy_migration(archive: ZipFile, target: str, column: int) -> None:
    """Refuse moving a legacy note when opaque Excel references may depend on it.

    The ordinary previous-version export has plain cells and a trailing note.
    Translating arbitrary formulas, table metadata or drawing anchors would need
    to rewrite other package parts, so those workbooks require explicit editing.
    """
    reason = "旧版 note 列包含复杂引用，无法安全移动；请先手动在 note 前插入 name，并在 note 后添加 not_bought、not_welded"
    opaque_parts = ("xl/tables/", "xl/drawings/", "xl/charts/", "xl/pivot", "xl/comments", "xl/threadedComments/", "xl/externalLinks/")
    for name in archive.namelist():
        if name.startswith(opaque_parts) or name.endswith(".vml"):
            raise BomError(reason)
        if not name.startswith("xl/") or not name.endswith(".xml"):
            continue
        root = ET.fromstring(archive.read(name))
        for node in root.iter():
            local = node.tag.rsplit("}", 1)[-1]
            if local in {"f", "formula", "formula1", "formula2"}:
                raise BomError(reason)
            if local == "definedName":
                if node.get("name") not in {"_xlnm.Print_Area", "_xlnm.Print_Titles"}:
                    raise BomError(reason)
                for area in (node.text or "").split(","):
                    if _range_reaches_column(area.rsplit("!", 1)[-1], column):
                        raise BomError(reason)
            if name == target and local not in {"c", "dimension"}:
                for attribute in ("ref", "sqref"):
                    if node.get(attribute) and _range_reaches_column(node.get(attribute), column):
                        raise BomError(reason)
                if local == "col" and int(node.get("max", "0")) > column:
                    raise BomError(reason)


def _move_legacy_note(xml: bytes, column: int) -> bytes:
    """Move the last plain note column by one, retaining all its cell content."""
    with minidom.parseString(xml) as document:
        root = document.documentElement
        data = _children(root, "sheetData")[0]
        for row in _children(data, "row"):
            for cell in _children(row, "c"):
                left, number, _, _ = range_boundaries(cell.getAttribute("r"))
                if left == column:
                    cell.setAttribute("r", f"{get_column_letter(column + 1)}{number}")
                    if row.hasAttribute("spans"):
                        bounds = [int(part) for span in row.getAttribute("spans").split() for part in span.split(":")]
                        row.setAttribute("spans", f"{min(bounds)}:{max(*bounds, column + 1)}")
        for cols in _children(root, "cols"):
            for item in list(_children(cols, "col")):
                left, right = int(item.getAttribute("min")), int(item.getAttribute("max"))
                if left <= column <= right:
                    moved = item.cloneNode(deep=True)
                    moved.setAttribute("min", str(column + 1))
                    moved.setAttribute("max", str(column + 1))
                    cols.insertBefore(moved, item.nextSibling)
                    if left < column:
                        item.setAttribute("max", str(column - 1))
                    else:
                        cols.removeChild(item)
                        item.unlink()
        return document.toxml(encoding="UTF-8", standalone=document.standalone)


def update_review_columns(source: bytes, sheet: str, updates: dict[int, dict[int, str]],
                          reference_column: int, *, widths: dict[int, int],
                          legacy_note: int | None = None) -> bytes:
    if not any(updates.values()):
        return source
    result = BytesIO()
    with ZipFile(BytesIO(source)) as original, ZipFile(result, "w") as output:
        target = _worksheet_part(original, sheet)
        if legacy_note is not None:
            _check_legacy_migration(original, target, legacy_note)
        for entry in original.infolist():
            data = original.read(entry.filename)
            if entry.filename == target:
                if legacy_note is not None:
                    data = _move_legacy_note(data, legacy_note)
                for column, values in sorted(updates.items()):
                    if values:
                        data = _patch_sheet(data, column, values, reference_column, widths.get(column))
            output.writestr(entry, data)
        output.comment = original.comment
    return result.getvalue()
